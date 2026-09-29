#!/usr/bin/env python3
"""
BIVA · lee la planilla de productos y actualiza products.json.

    python3 _tools/sync_sheet.py                      # usa la variable SHEET_CSV_URL
    python3 _tools/sync_sheet.py "https://docs.google.com/…/pub?output=csv"
    python3 _tools/sync_sheet.py productos.xlsx       # o un archivo local (.xlsx / .csv)

- Una fila = un producto. La columna SKU manda.
- La dirección web (slug) de cada SKU se conserva aunque cambies el nombre.
- Filas sin SKU o sin nombre se ignoran; SKU repetidos también (se avisa).
Después correr  python3 _tools/build.py  (el workflow de GitHub lo hace solo).
"""
import csv, io, json, os, re, sys, pathlib, unicodedata, urllib.request

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from sheet_schema import COLUMNS, CATEGORIES, AVAILABILITY, IMAGE_FIT

ROOT = pathlib.Path(__file__).resolve().parent.parent


def norm(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def slugify(s):
    return re.sub(r'-+', '-', re.sub(r'[^a-z0-9]+', '-', norm(s))).strip('-')


def split_list(v, lines=False):
    v = str(v or '').strip()
    if not v:
        return []
    parts = re.split(r'\r?\n|\||;', v) if lines else re.split(r',|\r?\n|;', v)
    return [p.strip(' -•\t') for p in parts if p.strip(' -•\t')]


def read_rows(source):
    if re.match(r'https?://', source):
        req = urllib.request.Request(source, headers={'User-Agent': 'biva-sync'})
        with urllib.request.urlopen(req, timeout=60) as r:
            raw = r.read().decode('utf-8-sig')
        return list(csv.reader(io.StringIO(raw)))
    path = pathlib.Path(source)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb['Productos'] if 'Productos' in wb.sheetnames else wb.worksheets[0]
        return [[('' if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
    return list(csv.reader(io.StringIO(path.read_text(encoding='utf-8-sig'))))


def lookup(mapping, value, default):
    n = norm(value)
    for key, label in mapping.items():
        if n and (n in (norm(label), norm(key)) or norm(label).startswith(n)):
            return key
    return default


def main():
    source = sys.argv[1] if len(sys.argv) > 1 else os.environ.get('SHEET_CSV_URL', '').strip()
    if not source:
        print('Sin planilla configurada (SHEET_CSV_URL vacío): se usa products.json tal cual.')
        return 0

    rows = read_rows(source)
    # find header row (first row that contains "SKU")
    hdr_i = next((i for i, r in enumerate(rows) if any(norm(c) == 'sku' for c in r)), None)
    if hdr_i is None:
        print('✗ No encontré la fila de encabezados (con la columna "SKU"). No se cambió nada.')
        return 1
    header = [norm(c) for c in rows[hdr_i]]
    col = {}
    for label, key, *_ in COLUMNS:
        if norm(label) in header:
            col[key] = header.index(norm(label))
    for req in ('sku', 'name'):
        if req not in col:
            print(f'✗ Falta la columna obligatoria "{req}". No se cambió nada.')
            return 1

    pj = ROOT / 'products.json'
    data = json.loads(pj.read_text(encoding='utf-8'))
    prev = {p['sku']: p for p in data['products']}
    used_slugs = {p['slug'] for p in data['products'] if p.get('slug')}

    out, seen, warnings = [], set(), []
    for n, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        get = lambda k: (str(r[col[k]]).strip() if k in col and col[k] < len(r) and r[col[k]] is not None else '')
        sku, name = get('sku'), get('name')
        if not sku and not name:
            continue
        if not sku or not name:
            warnings.append(f'Fila {n}: falta SKU o nombre, se ignoró.')
            continue
        if sku in seen:
            warnings.append(f'Fila {n}: SKU {sku} repetido, se ignoró.')
            continue
        seen.add(sku)
        old = prev.get(sku, {})

        slug = old.get('slug') or get('slug') or slugify(name)
        if not old.get('slug'):
            base, i = slug, 2
            while slug in used_slugs:
                slug, i = f'{base}-{i}', i + 1
            used_slugs.add(slug)

        cat_raw = get('category')
        category = lookup(CATEGORIES, cat_raw, None) or slugify(cat_raw) or old.get('category', 'cocina')
        if category not in CATEGORIES:
            warnings.append(f'Fila {n} ({sku}): categoría "{cat_raw}" desconocida; por ahora solo Cocina o Jardín.')

        p = {
            'id': old.get('id') or slug,
            'sku': sku,
            'slug': slug,
            'name': name,
            'shortDescription': get('shortDescription'),
            'description': get('description') or get('shortDescription'),
            'category': category,
            'tags': split_list(get('tags')),
            'availability': lookup(AVAILABILITY, get('availability'), 'in_stock'),
            'published': norm(get('published')) not in ('no', 'n', 'false', '0', 'oculto'),
            'imageFit': lookup(IMAGE_FIT, get('imageFit'), 'cover'),
            'wholesaleOnly': True,
            'priceOnRequest': True,
            'useCases': split_list(get('useCases')),
            'materials': split_list(get('materials')),
            'features': split_list(get('features'), lines=True),
        }
        for k in ('packSize', 'minOrder', 'dimensions'):
            if get(k):
                p[k] = get(k)
        # keep generated fields until build.py refreshes them
        for k in ('gallery', 'thumb', 'image', 'url'):
            if k in old:
                p[k] = old[k]
        out.append(p)

    if not out:
        print('✗ La planilla no tiene productos válidos. No se cambió nada (para no vaciar la web).')
        return 1

    removed = [s for s in prev if s not in seen]
    data['products'] = out
    pj.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    pub = sum(p['published'] for p in out)
    summary = [f'Planilla · {len(out)} productos ({pub} publicados, {len(out) - pub} ocultos)']
    if removed:
        summary.append(f'  Quitados de la planilla: {", ".join(removed)}')
    summary += [f'  ⚠ {w}' for w in warnings]
    print('\n'.join(summary))
    if os.environ.get('GITHUB_STEP_SUMMARY'):
        with open(os.environ['GITHUB_STEP_SUMMARY'], 'a', encoding='utf-8') as f:
            f.write('\n'.join(['### Planilla de productos', ''] + [f'- {s.strip()}' for s in summary]) + '\n')
    return 0


if __name__ == '__main__':
    sys.exit(main())
