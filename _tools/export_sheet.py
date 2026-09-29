#!/usr/bin/env python3
"""
BIVA · exporta products.json a una planilla Excel lista para usar
(o para importar en Google Sheets).

    python3 _tools/export_sheet.py [salida.xlsx]
"""
import json, sys, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from sheet_schema import COLUMNS, CATEGORIES, AVAILABILITY, IMAGE_FIT

ROOT = pathlib.Path(__file__).resolve().parent.parent
out = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'BIVA-productos.xlsx'
data = json.loads((ROOT / 'products.json').read_text(encoding='utf-8'))

CARBON, SALVIA, CREMA, PALE = '333330', '6B7A5C', 'EEEAE2', 'F0F4EC'
F = 'Arial'
thin = Side(style='thin', color='D9D6CE')

wb = Workbook()
ws = wb.active
ws.title = 'Productos'

for c, (label, key, width, help_) in enumerate(COLUMNS, 1):
    cell = ws.cell(row=1, column=c, value=label)
    cell.font = Font(name=F, bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor=CARBON)
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    cell.comment = Comment(help_, 'BIVA')
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'C2'

def value(p, key):
    v = p.get(key, '')
    if key == 'category':
        return CATEGORIES.get(v, v)
    if key == 'availability':
        return AVAILABILITY.get(v, 'En stock')
    if key == 'published':
        return 'Sí' if p.get('published', True) else 'No'
    if key == 'imageFit':
        return IMAGE_FIT.get(v, 'Recortada')
    if key == 'features':
        return '\n'.join(v or [])
    if isinstance(v, list):
        return ', '.join(v)
    return v

for r, p in enumerate(data['products'], 2):
    for c, (label, key, *_rest) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=r, column=c, value=value(p, key))
        cell.font = Font(name=F, size=10, color='787870' if key == 'slug' else CARBON)
        cell.alignment = Alignment(vertical='top', wrap_text=key in ('description', 'features', 'shortDescription'))
        cell.border = Border(bottom=thin)
        if key in ('availability', 'published'):
            cell.fill = PatternFill('solid', fgColor=PALE)

last = 2000
def dropdown(key, options):
    idx = [k for _, k, *_ in COLUMNS].index(key) + 1
    letter = ws.cell(row=1, column=idx).column_letter
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Valor no válido',
                        error='Elegí una opción de la lista.')
    dv.add(f'{letter}2:{letter}{last}')
    ws.add_data_validation(dv)

dropdown('category', list(CATEGORIES.values()))
dropdown('availability', list(AVAILABILITY.values()))
dropdown('published', ['Sí', 'No'])
dropdown('imageFit', list(IMAGE_FIT.values()))
ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{len(data["products"]) + 1}'

# --- instrucciones ---
g = wb.create_sheet('Cómo usar')
g.column_dimensions['A'].width = 26
g.column_dimensions['B'].width = 90
g['A1'] = 'Planilla de productos · bivasil.com'
g['A1'].font = Font(name=F, bold=True, size=14, color=CARBON)
rows = [
    ('', ''),
    ('Qué es', 'Esta planilla es la fuente de la web. Cada fila es un producto; la web se actualiza sola a partir de acá.'),
    ('Dónde se edita', 'En la pestaña "Productos". Las celdas verdes claras (Estado y Publicado) son las que más vas a tocar.'),
    ('Marcar sin stock', 'Columna Estado → "Sin stock". La página sigue online pero muestra "Sin stock" y el botón pide aviso de reposición.'),
    ('Ocultar un producto', 'Columna Publicado → "No". Desaparece del catálogo y su página deja de existir.'),
    ('Producto nuevo', '1) Agregá una fila con SKU nuevo, nombre, categoría y descripción. 2) Subí las fotos a la carpeta _fotos/<SKU>/ del repositorio (1.jpg es la principal). Sin fotos no se publica.'),
    ('Fotos', 'Van en GitHub, carpeta _fotos/<SKU>/, numeradas 1, 2, 3… Cualquier tamaño: la web genera versiones livianas.'),
    ('Listas', 'Materiales, Ideal para y Etiquetas: separados por coma. Características: una por línea (Alt+Enter en Excel, Ctrl+Enter en Google Sheets).'),
    ('No cambiar', 'SKU y Dirección web de productos ya publicados: cambiarlos rompe los links compartidos y los resultados de Google.'),
    ('Cuándo se ve', 'La web revisa la planilla cada hora. Para publicar ya: GitHub → Actions → "Actualizar productos" → Run workflow.'),
    ('Ayuda por columna', 'Pasá el mouse sobre cada encabezado de "Productos": tiene una nota con lo que va en esa columna.'),
]
for i, (a, b) in enumerate(rows, 2):
    g.cell(row=i, column=1, value=a).font = Font(name=F, bold=True, size=10, color=SALVIA)
    c = g.cell(row=i, column=2, value=b)
    c.font = Font(name=F, size=10, color=CARBON)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    g.row_dimensions[i].height = 30 if len(b) > 95 else 16

wb.save(out)
print(f'OK · {len(data["products"])} productos → {out}')
